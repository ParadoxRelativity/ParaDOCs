"""
Builds powerpoint-quarterly.pptx with python-pptx on its stock (PowerPoint-authored) template.

Run:  python3 -m venv /tmp/v && /tmp/v/bin/pip install python-pptx && /tmp/v/bin/python build_powerpoint.py

The deck exercises: title and content layouts with levels, a bulleted and a
numbered list, a hyperlink, formatted runs, a picture (generated PNG, cropped),
a table with a merged header and cell fills, preset shapes and a connector, a
group, speaker notes, a hidden slide, a transition and click animations.
"""
import io, struct, sys, zlib
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree


def png(w, h, rgb):
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))

    def chunk(t, d):
        c = struct.pack('>I', len(d)) + t + d
        return c + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)

    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')


prs = Presentation()
L = prs.slide_layouts

# 1 title
s = prs.slides.add_slide(L[0])
s.shapes.title.text = 'Quarterly Review'
s.placeholders[1].text = 'Finance and Operations'
s.notes_slide.notes_text_frame.text = 'Welcome everyone.\nKeep this to two minutes.'

# 2 bullets with levels, formatting, link
s = prs.slides.add_slide(L[1])
s.shapes.title.text = 'Highlights'
tf = s.placeholders[1].text_frame
tf.text = 'Revenue up 12%'
p = tf.add_paragraph()
p.text = 'Driven by new accounts'
p.level = 1
p = tf.add_paragraph()
p.level = 0
r = p.add_run(); r.text = 'Bold '; r.font.bold = True
r = p.add_run(); r.text = 'italic '; r.font.italic = True
r = p.add_run(); r.text = 'red'; r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00); r.font.size = Pt(28)
r = p.add_run(); r.text = ' link'; r.hyperlink.address = 'https://example.com/report'
s.notes_slide.notes_text_frame.text = 'Mention the churn numbers.'

# 3 picture + table + shapes
s = prs.slides.add_slide(L[5])
s.shapes.title.text = 'Pipeline'
pic = s.shapes.add_picture(io.BytesIO(png(64, 48, (30, 120, 200))), Inches(0.6), Inches(1.6), Inches(3), Inches(2.25))
pic.crop_left = 0.1
pic.crop_right = 0.1
gf = s.shapes.add_table(3, 3, Inches(3.8), Inches(1.6), Inches(5.6), Inches(2))
t = gf.table
t.cell(0, 0).merge(t.cell(0, 1))
t.cell(0, 0).text = 'Region'
t.cell(0, 2).text = 'Total'
t.cell(1, 0).text = 'North'; t.cell(1, 1).text = '12'; t.cell(1, 2).text = '30'
t.cell(2, 0).text = 'South'; t.cell(2, 1).text = '8'; t.cell(2, 2).text = '22'
t.cell(1, 2).fill.solid()
t.cell(1, 2).fill.fore_color.rgb = RGBColor(0xFF, 0xF2, 0xCC)
box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(4.5), Inches(3), Inches(1.2))
box.fill.solid()
box.fill.fore_color.rgb = RGBColor(0x2E, 0x75, 0xB6)
box.text_frame.text = 'On track'
box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(3.8), Inches(4.5), Inches(1.8), Inches(1.2))
s.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(5.8), Inches(4.5), Inches(1.4), Inches(1.2))
s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(7.4), Inches(4.5), Inches(9.4), Inches(5.7))
grp = s.shapes.add_group_shape()
grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(7.8), Inches(3.8), Inches(0.8), Inches(0.8))
grp.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.8), Inches(3.8), Inches(0.8), Inches(0.8))

# 4 hidden numbered list and a free text box
s = prs.slides.add_slide(L[1])
s.shapes.title.text = 'Appendix (hidden)'
tf = s.placeholders[1].text_frame
tf.text = 'First'
tf.add_paragraph().text = 'Second'
tb = s.shapes.add_textbox(Inches(1), Inches(6.4), Inches(6), Inches(0.6))
tb.text_frame.text = 'Source: internal data'
tb.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
s._element.set('show', '0')
for para in tf.paragraphs:
    pPr = para._p.get_or_add_pPr()
    for old in pPr.findall(qn('a:buNone')):
        pPr.remove(old)
    etree.SubElement(pPr, qn('a:buAutoNum')).set('type', 'arabicPeriod')

# 5 closing slide with a transition and click animations
s = prs.slides.add_slide(L[2])
s.shapes.title.text = 'Questions'
s.placeholders[1].text = 'Thank you'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
tr = etree.SubElement(s._element, '{%s}transition' % P)
tr.set('spd', 'med')
etree.SubElement(tr, '{%s}fade' % P)
sid = s.shapes.title.shape_id
sid2 = s.placeholders[1].shape_id


def effect(i, spid, preset, subtype, node):
    return f'''<p:par><p:cTn id="{i}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="{i+1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:par><p:cTn id="{i+2}" presetID="{preset}" presetClass="entr" presetSubtype="{subtype}" fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
<p:set><p:cBhvr><p:cTn id="{i+3}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>
<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{i+4}" dur="500"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>
</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'''


timing = f'''<p:timing xmlns:p="{P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{effect(3, sid, 10, 0, 'clickEffect')}{effect(10, sid2, 2, 4, 'clickEffect')}</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''
s._element.append(etree.fromstring(timing))
s.notes_slide.notes_text_frame.text = 'Open the floor.'

prs.core_properties.title = 'Quarterly Review'
prs.save(sys.argv[1] if len(sys.argv) > 1 else 'powerpoint-quarterly.pptx')
